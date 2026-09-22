"""Regenerates the stress-test PDF fixtures used to prove post-processing
fixes real Marker defects, not just hand-crafted markdown strings. See
tests/test_postprocess_integration.py for stress.pdf, and
tests/test_postprocess_more_integration.py for the ones generated here.

    python tests/fixtures/generate_stress_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

FIXTURES_DIR = Path(__file__).parent
styles = getSampleStyleSheet()
body = styles["BodyText"]
h_title = ParagraphStyle("H0", fontSize=28, leading=32, fontName="Helvetica-Bold")
h_big = ParagraphStyle("H1", fontSize=22, leading=26, fontName="Helvetica-Bold")


def generate_running_header() -> None:
    """4-page document with a bold header drawn at a fixed position on every
    page (not a flowable heading) - tests whether that gets misdetected as a
    duplicate heading, and whether structurally identical section headings
    get assigned consistent levels throughout."""

    def draw_header(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 14)
        canvas.drawString(72, 740, "CONFIDENTIAL - INTERNAL REPORT")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(FIXTURES_DIR / "running_header.pdf"), pagesize=LETTER)
    story = [Paragraph("Quarterly Risk Review", h_title), Spacer(1, 20)]
    for section in range(1, 5):
        story.append(Paragraph(f"Section {section}: Findings", h_big))
        story.append(
            Paragraph(
                (f"This is body paragraph content for section {section}. ") * 20,
                body,
            )
        )
        story.append(PageBreak())
    doc.build(story, onFirstPage=draw_header, onLaterPages=draw_header)


def generate_merged_cells() -> None:
    """A table with a spanned header cell (merged columns) - tests whether
    that breaks raw column-count reconstruction."""
    rows = [
        ["", "Q1 2026", "", "Q2 2026", ""],
        ["Region", "Revenue", "Growth", "Revenue", "Growth"],
        ["West", "$1.2M", "+5%", "$1.4M", "+17%"],
        ["East", "$0.9M", "+2%", "$1.0M", "+11%"],
    ]
    doc = SimpleDocTemplate(str(FIXTURES_DIR / "merged_cells.pdf"), pagesize=LETTER)
    table = Table(
        rows,
        style=TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("SPAN", (1, 0), (2, 0)),
                ("SPAN", (3, 0), (4, 0)),
                ("BACKGROUND", (0, 0), (-1, 1), colors.lightgrey),
            ]
        ),
    )
    doc.build([Paragraph("Regional Revenue", h_big), Spacer(1, 10), table])


def generate_with_image() -> None:
    """A PDF with a real embedded image, to test image_handling.py against
    real Marker image extraction rather than a synthetic dict."""
    img_path = FIXTURES_DIR / "_gen_chart.png"
    PILImage.new("RGB", (300, 200), color=(70, 130, 180)).save(img_path)

    doc = SimpleDocTemplate(str(FIXTURES_DIR / "with_image.pdf"), pagesize=LETTER)
    doc.build(
        [
            Paragraph("Revenue Chart", h_big),
            Spacer(1, 10),
            Paragraph("See the chart below for regional performance.", body),
            Spacer(1, 10),
            Image(str(img_path), width=3 * inch, height=2 * inch),
            Spacer(1, 10),
            Paragraph("Chart data is preliminary and subject to revision.", body),
        ]
    )
    img_path.unlink()  # embedded in the PDF now, don't need the source file


def generate_rotated_page() -> None:
    """A 2-page document with the second page's `/Rotate` flag set to 90deg -
    tests whether the backend respects PDF rotation metadata rather than
    extracting the page's text sideways/garbled. Found worth testing after a
    real downloaded patent PDF with a rotated page converted correctly; this
    fixture keeps that regression covered without redistributing the real
    file (per this project's policy against committing real-world
    documents).

    Each page carries a full paragraph, not a single short sentence: a first
    attempt with one bare line per page passed locally but failed in Linux CI
    - the layout model classified the rotated page's lone line as a
    PageFooter (which Marker excludes from output) rather than body text, a
    borderline call that came out differently across platforms. That's the
    same backend-non-determinism class CONTRACT.md already documents for OCR;
    a real paragraph's worth of body text keeps this test clear of that
    boundary instead of asserting on it."""
    import pypdfium2 as pdfium

    doc = SimpleDocTemplate(str(FIXTURES_DIR / "_gen_rotated.pdf"), pagesize=LETTER)
    doc.build(
        [
            Paragraph(
                "This is the first page, and it is not rotated. " * 8,
                body,
            ),
            PageBreak(),
            Paragraph(
                "This is the second page, and its /Rotate flag is set to ninety "
                "degrees in the PDF itself, though the text should still read "
                "upright and in order once extracted. " * 6,
                body,
            ),
        ]
    )

    pdf = pdfium.PdfDocument(str(FIXTURES_DIR / "_gen_rotated.pdf"))
    pdf[1].set_rotation(90)
    pdf.save(str(FIXTURES_DIR / "rotated_page.pdf"))
    (FIXTURES_DIR / "_gen_rotated.pdf").unlink()


def generate_many_slides_pptx() -> None:
    """20 slides, each with enough body text that several land on the same
    rendered PDF page - reproduces a real Marker bug: IgnoreTextProcessor
    strips trailing digits before comparing text across pages, so docmd's
    own "Slide 1", "Slide 2", ... headings all collapse to the identical
    string "Slide" and get flagged as a repeated running header/footer,
    exactly like the heuristic they're meant to catch. See
    marker_converter._unsuppress_slide_headings()."""
    from pptx import Presentation

    prs = Presentation()
    for i in range(1, 21):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Topic number {i}"
        text_frame = slide.placeholders[1].text_frame
        text_frame.text = f"First point about topic {i}"
        for j in range(2, 6):
            text_frame.add_paragraph().text = f"Point {j} about topic {i} with some extra words to take space"
    prs.save(str(FIXTURES_DIR / "many_slides.pptx"))


if __name__ == "__main__":
    generate_running_header()
    generate_merged_cells()
    generate_with_image()
    generate_rotated_page()
    generate_many_slides_pptx()
    print(f"wrote stress fixtures to {FIXTURES_DIR}")
